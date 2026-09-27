from __future__ import annotations

import asyncio
import hashlib
import shutil
import subprocess
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from fastapi import UploadFile
from PIL import Image
from docx import Document

from app.imports.schemas import BoundingBox


class ImportStorageError(ValueError):
    pass


@dataclass(frozen=True)
class StoredUpload:
    original_name: str
    mime_type: str
    size: int
    sha256: str
    relative_path: str


@dataclass(frozen=True)
class CroppedAsset:
    relative_path: str
    sha256: str
    width: int
    height: int


class LocalImportStorage:
    def __init__(
        self,
        root: str | Path,
        *,
        max_file_size_mb: int = 50,
        render_dpi: int = 200,
        libreoffice_path: str = "soffice",
        use_local_ocr: bool = True,
    ) -> None:
        self.root = Path(root).resolve()
        self.max_bytes = max_file_size_mb * 1024 * 1024
        self.render_dpi = render_dpi
        self.libreoffice_path = libreoffice_path
        self.use_local_ocr = use_local_ocr

    def resolve(self, relative_path: str) -> Path:
        candidate = (self.root / relative_path).resolve()
        if candidate != self.root and self.root not in candidate.parents:
            raise ImportStorageError("stored path escapes import storage root")
        return candidate

    async def save_upload(self, batch_id: uuid.UUID, upload: UploadFile) -> StoredUpload:
        original_name = Path(upload.filename or "upload").name
        suffix = Path(original_name).suffix.lower()
        if suffix not in {".pdf", ".docx"}:
            raise ImportStorageError("only PDF and DOCX files are accepted")

        target_dir = self.resolve(f"{batch_id}/source")
        target_dir.mkdir(parents=True, exist_ok=True)
        temp_path = target_dir / f"upload-{uuid.uuid4().hex}.tmp"
        digest = hashlib.sha256()
        size = 0
        try:
            with temp_path.open("wb") as stream:
                while chunk := await upload.read(1024 * 1024):
                    size += len(chunk)
                    if size > self.max_bytes:
                        raise ImportStorageError(
                            f"file exceeds configured {self.max_bytes // (1024 * 1024)} MB limit"
                        )
                    digest.update(chunk)
                    stream.write(chunk)
            self._validate_signature(temp_path, suffix)
            sha256 = digest.hexdigest()
            final_path = target_dir / f"{sha256}{suffix}"
            if final_path.exists():
                temp_path.unlink()
            else:
                temp_path.replace(final_path)
        except Exception:
            temp_path.unlink(missing_ok=True)
            raise

        mime_type = (
            "application/pdf"
            if suffix == ".pdf"
            else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        return StoredUpload(
            original_name=original_name,
            mime_type=mime_type,
            size=size,
            sha256=sha256,
            relative_path=final_path.relative_to(self.root).as_posix(),
        )

    def _validate_signature(self, path: Path, suffix: str) -> None:
        if suffix == ".pdf":
            with path.open("rb") as stream:
                signature = stream.read(5)
            if signature != b"%PDF-":
                raise ImportStorageError("uploaded file is not a valid PDF")
            return
        if not zipfile.is_zipfile(path):
            raise ImportStorageError("uploaded file is not a valid DOCX archive")
        with zipfile.ZipFile(path) as archive:
            if "word/document.xml" not in archive.namelist():
                raise ImportStorageError("uploaded archive does not contain a Word document")

    def remove_batch(self, batch_id: uuid.UUID) -> None:
        batch_dir = self.resolve(str(batch_id))
        if batch_dir.is_dir():
            shutil.rmtree(batch_dir)

    async def render_document(
        self, batch_id: uuid.UUID, document_id: uuid.UUID, relative_source_path: str
    ) -> int:
        return await asyncio.to_thread(
            self._render_document_sync, batch_id, document_id, relative_source_path
        )

    async def extract_docx_content(
        self, batch_id: uuid.UUID, document_id: uuid.UUID, relative_source_path: str
    ) -> dict:
        return await asyncio.to_thread(
            self._extract_docx_content_sync,
            batch_id,
            document_id,
            relative_source_path,
        )

    def _extract_docx_content_sync(
        self, batch_id: uuid.UUID, document_id: uuid.UUID, relative_source_path: str
    ) -> dict:
        source_path = self.resolve(relative_source_path)
        if source_path.suffix.lower() != ".docx":
            return {}
        document = Document(source_path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        tables = [
            [[cell.text for cell in row.cells] for row in table.rows]
            for table in document.tables
        ]
        image_dir = self.resolve(f"{batch_id}/docx-assets/{document_id}")
        image_dir.mkdir(parents=True, exist_ok=True)
        images: list[dict[str, str]] = []
        seen_parts: set[str] = set()
        for relationship in document.part.rels.values():
            target = getattr(relationship, "target_part", None)
            partname = str(getattr(target, "partname", ""))
            if not partname.startswith("/word/media/") or partname in seen_parts:
                continue
            seen_parts.add(partname)
            filename = Path(partname).name
            output = image_dir / filename
            output.write_bytes(target.blob)
            images.append(
                {
                    "name": filename,
                    "content_type": target.content_type,
                    "storage_path": output.relative_to(self.root).as_posix(),
                }
            )
        return {"paragraphs": paragraphs, "tables": tables, "embedded_images": images}

    def _render_document_sync(
        self, batch_id: uuid.UUID, document_id: uuid.UUID, relative_source_path: str
    ) -> int:
        source_path = self.resolve(relative_source_path)
        render_dir = self.resolve(f"{batch_id}/pages/{document_id}")
        render_dir.mkdir(parents=True, exist_ok=True)
        pdf_path = source_path
        if source_path.suffix.lower() == ".docx":
            pdf_path = self._convert_docx_to_pdf(source_path, render_dir)

        document = pymupdf.open(pdf_path)
        try:
            if document.page_count < 1:
                raise ImportStorageError("document has no pages")
            scale = self.render_dpi / 72
            matrix = pymupdf.Matrix(scale, scale)
            for page_index in range(document.page_count):
                output = render_dir / f"page-{page_index + 1:04d}.png"
                if output.exists():
                    continue
                pixmap = document[page_index].get_pixmap(matrix=matrix, alpha=False)
                pixmap.save(output)
            return document.page_count
        finally:
            document.close()

    def _convert_docx_to_pdf(self, source_path: Path, render_dir: Path) -> Path:
        executable = shutil.which(self.libreoffice_path)
        if executable is None:
            common_path = Path(r"D:\Program Files\LibreOffice\program\soffice.com")
            if common_path.exists():
                executable = str(common_path)
        if executable is None:
            raise ImportStorageError("LibreOffice is required to preview DOCX files")
        result = subprocess.run(
            [
                executable,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(render_dir),
                str(source_path),
            ],
            check=False,
            capture_output=True,
            text=True,
            timeout=120,
        )
        output = render_dir / f"{source_path.stem}.pdf"
        if result.returncode != 0 or not output.exists():
            detail = (result.stderr or result.stdout or "unknown LibreOffice error").strip()
            raise ImportStorageError(f"DOCX preview conversion failed: {detail}")
        return output

    def page_image_path(
        self, batch_id: uuid.UUID, document_id: uuid.UUID, page_no: int
    ) -> Path:
        path = self.resolve(f"{batch_id}/pages/{document_id}/page-{page_no:04d}.png")
        if not path.is_file():
            raise ImportStorageError(f"rendered page {page_no} does not exist")
        return path

    def extraction_cache_path(
        self, batch_id: uuid.UUID, document_id: uuid.UUID, page_no: int
    ) -> Path:
        return self.resolve(f"{batch_id}/extraction/{document_id}/page-{page_no:04d}.json")

    async def crop_asset(
        self,
        batch_id: uuid.UUID,
        document_id: uuid.UUID,
        page_no: int,
        bbox: BoundingBox,
        asset_id: uuid.UUID,
    ) -> CroppedAsset:
        return await asyncio.to_thread(
            self._crop_asset_sync, batch_id, document_id, page_no, bbox, asset_id
        )

    def _crop_asset_sync(
        self,
        batch_id: uuid.UUID,
        document_id: uuid.UUID,
        page_no: int,
        bbox: BoundingBox,
        asset_id: uuid.UUID,
    ) -> CroppedAsset:
        page_path = self.page_image_path(batch_id, document_id, page_no)
        with Image.open(page_path) as source:
            width, height = source.size
            box = (
                round(bbox.x0 * width),
                round(bbox.y0 * height),
                round(bbox.x1 * width),
                round(bbox.y1 * height),
            )
            cropped = source.crop(box).convert("RGB")
            asset_dir = self.resolve(f"{batch_id}/assets")
            asset_dir.mkdir(parents=True, exist_ok=True)
            output = asset_dir / f"{asset_id}.png"
            cropped.save(output, format="PNG", optimize=True)
            asset_width, asset_height = cropped.size
        digest = hashlib.sha256(output.read_bytes()).hexdigest()
        return CroppedAsset(
            relative_path=output.relative_to(self.root).as_posix(),
            sha256=digest,
            width=asset_width,
            height=asset_height,
        )
