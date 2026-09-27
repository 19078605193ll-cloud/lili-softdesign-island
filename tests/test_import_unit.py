from decimal import Decimal
from pathlib import Path
import uuid
from unittest.mock import AsyncMock

import pytest
from docx import Document
from pydantic import ValidationError
from PIL import Image

from app.imports.ai import OpenAICompatibleQuestionClient

from app.imports.schemas import (
    BoundingBox,
    BulkApproveItemsInput,
    KnowledgeSelectionInput,
    SourceSectionSelection,
    VisionCombinedPage,
)
from app.imports.service import (
    _repair_pages_for_question,
    build_batch_validation_summary,
    infer_repeated_stem_ranges,
    parse_explicit_question_range,
    validate_import_item,
)
from app.imports.storage import LocalImportStorage
from app.models import QuestionImportItem


async def test_combined_page_uses_one_model_request_and_requires_both_roles(tmp_path: Path) -> None:
    with pytest.raises(ValidationError):
        VisionCombinedPage.model_validate({"groups": [], "answers": [{
            "question_no": 1, "correct_option_keys": ["A"]
        }]})
    payload = {
        "groups": [{"source_label": "1", "items": [{"question_no": 1, "stem_markdown": "题干"}]}],
        "answers": [{"question_no": 1, "correct_option_keys": ["A"], "explanation_markdown": "解析"}],
    }
    image = tmp_path / "page.png"
    Image.new("RGB", (10, 10), "white").save(image)
    client = object.__new__(OpenAICompatibleQuestionClient)
    client.vision_model = "test-vision"
    client._request_json = AsyncMock(return_value=(VisionCombinedPage.model_validate(payload), payload))
    result, raw = await client.parse_combined_page(image, 1)
    assert result.answers[0].question_no == 1
    assert raw == payload
    client._request_json.assert_awaited_once()
    assert client._request_json.await_args.kwargs["schema"] is VisionCombinedPage


def test_source_sections_allow_cross_role_overlap_but_not_same_role_overlap() -> None:
    selection = SourceSectionSelection.model_validate(
        {"questions": [{"start": 1, "end": 3}], "answers": [{"start": 3, "end": 5}]}
    )
    assert selection.questions[0].end == selection.answers[0].start
    with pytest.raises(ValidationError):
        SourceSectionSelection.model_validate(
            {
                "questions": [{"start": 1, "end": 3}, {"start": 3, "end": 4}],
                "answers": [{"start": 1, "end": 2}],
            }
        )


def test_bounding_box_requires_positive_area() -> None:
    with pytest.raises(ValidationError):
        BoundingBox(x0=0.5, y0=0.2, x1=0.4, y1=0.8)


def test_knowledge_selection_rejects_duplicate_or_primary_related_codes() -> None:
    with pytest.raises(ValidationError):
        KnowledgeSelectionInput(
            primary_code="topic.primary",
            related_codes=["topic.related", "topic.related"],
        )
    with pytest.raises(ValidationError):
        KnowledgeSelectionInput(
            primary_code="topic.primary",
            related_codes=["topic.primary"],
        )


def test_bulk_approval_requires_unique_item_ids() -> None:
    item_id = uuid.uuid4()
    with pytest.raises(ValidationError):
        BulkApproveItemsInput(item_ids=[item_id, item_id])


def test_item_validation_enforces_morning_question_shape() -> None:
    item = QuestionImportItem(
        batch_id="00000000-0000-0000-0000-000000000001",
        question_no=1,
        question_type="single_choice",
        stem_markdown="题干",
        options_payload=[
            {"key": "A", "content_markdown": "选项 A"},
            {"key": "B", "content_markdown": "选项 B"},
        ],
        correct_option_keys=["C"],
        score=Decimal("1"),
        source_refs=[],
    )
    issues = validate_import_item(item)
    codes = {issue.code for issue in issues}
    assert "invalid_options" in codes
    assert "answer_not_in_options" in codes
    assert "missing_explanation" in codes


def test_only_explicit_number_ranges_create_shared_groups() -> None:
    assert parse_explicit_question_range("42-43") == [42, 43]
    assert parse_explicit_question_range("71~75") == [71, 72, 73, 74, 75]
    assert parse_explicit_question_range("8") is None
    assert parse_explicit_question_range("page-8") is None


def test_batch_validation_detects_missing_numbers_and_structural_errors() -> None:
    items = [
        QuestionImportItem(
            batch_id="00000000-0000-0000-0000-000000000001",
            question_no=number,
            question_type="single_choice",
            stem_markdown=f"question {number}",
            options_payload=[
                {"key": key, "content_markdown": f"option {key}"}
                for key in ["A", "B", "C", "D"]
            ],
            correct_option_keys=["A"],
            score=Decimal("1"),
            source_refs=[],
            validation_issues=[],
            status="needs_review",
        )
        for number in range(1, 76)
        if number != 43
    ]
    item_42 = next(item for item in items if item.question_no == 42)
    item_42.options_payload = item_42.options_payload[:2]
    item_42.validation_issues = [
        {"code": "invalid_options", "severity": "error", "message": "invalid"}
    ]

    summary = build_batch_validation_summary(items, expected_count=75)

    assert summary["is_valid"] is False
    assert summary["missing_question_numbers"] == [43]
    assert summary["repairable_question_numbers"] == [42, 43]
    assert summary["blocking_issue_count"] == 2


def test_incomplete_item_prefers_following_page_for_continuation() -> None:
    refs = {
        54: {"questions": {9, 10}, "answers": set()},
        55: {"questions": {10}, "answers": set()},
        56: {"questions": {11}, "answers": set()},
    }

    pages = _repair_pages_for_question(
        55,
        role="questions",
        refs_by_question=refs,
        selected_pages=list(range(1, 14)),
    )

    assert pages == [10, 11]


def test_consecutive_items_with_identical_stems_form_shared_range() -> None:
    items_data = {
        67: {"stem": "ordinary question", "group_label": "single-67"},
        68: {"stem": "shared prompt with two independent blanks", "group_label": "single-68"},
        69: {"stem": "shared prompt with two independent blanks", "group_label": "single-69"},
        70: {"stem": "another ordinary question", "group_label": "single-70"},
    }

    assert infer_repeated_stem_ranges(items_data) == [
        (68, 69, "shared prompt with two independent blanks")
    ]


async def test_docx_direct_extraction_preserves_paragraphs_and_tables(tmp_path: Path) -> None:
    storage = LocalImportStorage(tmp_path / "imports")
    batch_id = uuid.uuid4()
    document_id = uuid.uuid4()
    relative_path = f"{batch_id}/source/sample.docx"
    source_path = storage.resolve(relative_path)
    source_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph("第 1 题题干")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    document.save(source_path)

    extracted = await storage.extract_docx_content(
        batch_id, document_id, relative_path
    )
    assert extracted["paragraphs"] == ["第 1 题题干"]
    assert extracted["tables"] == [[["A", "B"], ["1", "2"]]]
    assert extracted["embedded_images"] == []
